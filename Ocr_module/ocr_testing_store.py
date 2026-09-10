"""
Persistence for the PaddleOCR-vs-Tesseract engine testing UI (see
routes/testing_routes.py and the "OCR Engine Testing" panel in
templates/index.html).

Unlike case_session_store.py (one row per case, replaced on every save), this
keeps a single running workbook that every save *appends* to - so results
accumulate across row navigation, page navigation, and new file uploads
within the same testing session, matching the capstone's per-field
comparison sheet (Test ID, Patient ID, Field, Ground Truth, PaddleOCR
Output, Paddle Correct?, TesseractOCR Output, Tesseract Correct?).

File-based (openpyxl), not database-backed: there is no schema for this data
in the project's MariaDB instance (research_ground_truth/research_ocr_results
are documented in benchmarks/README.md as future intent only - see that file
for details), and an accumulating .xlsx is exactly the deliverable requested.
"""
import logging
import threading
from datetime import datetime

from openpyxl import Workbook, load_workbook

import config

logger = logging.getLogger(__name__)

HEADERS = [
    "Test ID",
    "Spread ID",
    "Source File",
    "PDF Page",
    "Patient ID",
    "Field",
    "Ground Truth",
    "PaddleOCR Output",
    "Paddle Confidence",
    "Paddle Correct?",
    "TesseractOCR Output",
    "Tesseract Confidence",
    "Tesseract Correct?",
    "Saved At",
]

SHEET_NAME = "OCR Testing"

# Guards the workbook file across concurrent requests (Flask's dev server can
# handle requests on multiple threads) - without this, two near-simultaneous
# saves could both read the same "current" row count and overwrite each other.
_lock = threading.Lock()


def _load_or_create_workbook():
    path = config.OCR_TESTING_WORKBOOK
    if path.exists():
        workbook = load_workbook(path)
        if SHEET_NAME in workbook.sheetnames:
            sheet = workbook[SHEET_NAME]
        else:
            sheet = workbook.create_sheet(SHEET_NAME)
            sheet.append(HEADERS)
        return workbook, sheet

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)
    return workbook, sheet


def _bool_label(value):
    if value is None or value == "":
        return ""
    return "Yes" if value in (True, "true", "True", "1", 1) else "No"


def append_rows(rows):
    """Append one row per item in `rows` to the running workbook.

    Each item is a dict with keys matching the (snake_case) fields the
    frontend collects per tested cell - see routes/testing_routes.py for the
    exact accepted keys. Returns {"added": int, "total_rows": int}.
    """
    if not rows:
        return {"added": 0, "total_rows": count_rows()}

    with _lock:
        workbook, sheet = _load_or_create_workbook()
        saved_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        added = 0
        for row in rows:
            # Header occupies row 1, so before inserting the Nth data row, max_row == N -
            # exactly the Test ID that row should get. Recomputed each iteration since
            # max_row grows by one with every append() below.
            next_test_id = sheet.max_row
            sheet.append([
                next_test_id,
                row.get("spread_id", ""),
                row.get("source_file", ""),
                row.get("pdf_page", ""),
                row.get("patient_id", ""),
                row.get("field", ""),
                row.get("ground_truth", ""),
                row.get("paddle_output", ""),
                row.get("paddle_confidence", ""),
                _bool_label(row.get("paddle_correct")),
                row.get("tesseract_output", ""),
                row.get("tesseract_confidence", ""),
                _bool_label(row.get("tesseract_correct")),
                saved_at,
            ])
            added += 1

        workbook.save(config.OCR_TESTING_WORKBOOK)
        total_rows = sheet.max_row - 1  # minus the header row

    logger.info("Appended %d OCR test row(s); workbook now has %d row(s)", added, total_rows)
    return {"added": added, "total_rows": total_rows}


def count_rows():
    with _lock:
        if not config.OCR_TESTING_WORKBOOK.exists():
            return 0
        workbook = load_workbook(config.OCR_TESTING_WORKBOOK, read_only=True)
        sheet = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook.active
        return max(sheet.max_row - 1, 0)


def clear_rows():
    """Wipe every saved test row, leaving a fresh workbook with just the header.

    Used by the "Clear Saved Rows" button in the testing panel - a deliberate,
    explicit reset (confirmed client-side before this is ever called), not
    something any other operation triggers as a side effect.
    """
    with _lock:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = SHEET_NAME
        sheet.append(HEADERS)
        workbook.save(config.OCR_TESTING_WORKBOOK)

    logger.info("Cleared all saved OCR test rows")
    return {"total_rows": 0}


def get_workbook_path():
    return config.OCR_TESTING_WORKBOOK
